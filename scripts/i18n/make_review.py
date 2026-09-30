"""Build docs/i18n-review/translation-review.xlsx from web/src/i18n/{en,hi,mr}.json
for a native-speaker review. Run: python scripts/i18n/make_review.py"""
import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

import os
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
base = os.path.join(ROOT, "web/src/i18n/")
en, hi, mr = (json.load(open(base + l + ".json", encoding="utf8")) for l in ("en", "hi", "mr"))

def flat(d, p=""):
    for k, v in d.items():
        kk = f"{p}.{k}" if p else k
        if isinstance(v, dict): yield from flat(v, kk)
        else: yield kk, v
def get(d, key):
    for part in key.split("."): d = d[part]
    return d

SCREEN = {"app":"App shell","nav":"Navigation","role":"Roles","offline":"Offline","bill":"Bill (counter)","pending":"Pending tokens",
 "pay":"Payment","receipt":"Receipt slip","unit":"Units","amend":"Edit bill","completed":"Completed bills","dues":"Dues (udhaar)",
 "close":"Close day","void":"Void bill","error":"Error messages","session":"Sign-in / session","changePw":"Change password",
 "items":"Items","stock":"Stock","customersScreen":"Customers","rateList":"Rate-list photo import","range":"Date filter",
 "dash":"Dashboards","staff":"Staff","settings":"Settings","shop":"Shop","danger":"Danger zone","outbox":"Outbox (offline queue)",
 "syncIssues":"Sync issues","req":"Requests","soon":"Coming soon","owner":"Platform owner console"}
HIGH = {"app","nav","offline","bill","pending","pay","receipt","unit","amend","error","dues","close","void","role"}
MED = {"completed","items","stock","customersScreen","session","changePw","outbox","syncIssues","range"}
def prio(g): return "1 - High" if g in HIGH else "2 - Medium" if g in MED else "3 - Low"

rows = []
for key, e in flat(en):
    g = key.split(".")[0]
    rows.append((prio(g), SCREEN.get(g, g), key, e, get(hi, key), get(mr, key)))
rows.sort(key=lambda r: (r[0], r[1], r[2]))

F = "Nirmala UI"
hdr_font = Font(name=F, bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor="047857")
inp = PatternFill("solid", fgColor="FFF59D"); thin = Side(style="thin", color="CBD5E1")
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
ins = wb.active; ins.title = "Instructions"
lines = [
 ("How to review", True),
 ("One tab per language: 'Marathi' and 'Hindi'. You only need the tab for your language.", False),
 ("Each row is one piece of text in the app. Read the English, then the current translation.", False),
 ("If the current text is fine, set Verdict to OK. If not, set Verdict to Change and type your version in the yellow 'Your version' cell.", False),
 ("Only edit the yellow columns (Verdict, Your version, Comment). Don't change other columns.", False),
 ("Start with Priority 1 rows: they are the screens staff use at the counter all day.", False),
 ("", False),
 ("Rules", True),
 ("Keep anything in double curly braces exactly as it is, e.g. {{count}}, {{date}}, {{name}}. The app fills these in.", False),
 ("Write how shop staff actually speak, not formal or textbook language. Short is better: people read this one-handed mid-sale.", False),
 ("Keys ending in _one / _other are singular / plural forms of the same message.", False),
 ("", False),
 ("Open questions (please answer in the Questions tab)", True),
 ("See the 'Questions' tab: a few word choices we could not decide without a native speaker.", False),
 ("", False),
 ("Example (Marathi)", True),
 ("English: 'Sign out'  |  Current: 'साइन आउट'  |  Verdict: OK", False),
 ("English: 'Close day'  |  Current: 'दिवस बंद करा'  |  Verdict: Change  |  Your version: 'दिवसाचा हिशोब बंद करा'  |  Comment: sounds more natural at a shop", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ins.cell(i, 1, t); c.font = Font(name=F, bold=b, size=12 if b else 11); c.alignment = Alignment(wrap_text=True)
ins.column_dimensions["A"].width = 120

cols = ["Priority", "Screen", "Key", "English", "Current", "Verdict", "Your version", "Comment"]
widths = [11, 22, 30, 45, 45, 11, 45, 30]
for name, idx in (("Marathi", 5), ("Hindi", 4)):
    ws = wb.create_sheet(name)
    for j, h in enumerate(cols, 1):
        c = ws.cell(1, j, h); c.font = hdr_font; c.fill = hdr_fill; c.alignment = wrap
        ws.column_dimensions[chr(64 + j)].width = widths[j - 1]
    dv = DataValidation(type="list", formula1='"OK,Change"', allow_blank=True); ws.add_data_validation(dv)
    for i, r in enumerate(rows, 2):
        vals = [r[0], r[1], r[2], r[3], r[idx], None, None, None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v); c.font = Font(name=F, size=11); c.alignment = wrap
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if j >= 6: c.fill = inp
        dv.add(f"F{i}")
    ws.freeze_panes = "D2"; ws.auto_filter.ref = f"A1:H{len(rows)+1}"

q = wb.create_sheet("Questions")
qs = [("Question", "Current choice", "Your answer (Marathi)", "Your answer (Hindi)"),
 ("Marathi sentences currently end with a full stop (.), Hindi with a danda (।). Which should Marathi use?", "Marathi: full stop", None, "n/a"),
 ("'टोकन' (token) is kept in English rather than translated. OK?", "Transliterated", None, None),
 ("'बास्केट' (basket) is kept in English rather than translated. OK?", "Transliterated", None, None),
 ("'पॉइंट्स' (loyalty points) is kept in English rather than translated. OK?", "Transliterated", None, None),
 ("Should the app speak to staff politely (आपण / आप) or casually (तू / तुम)?", "Mixed / unchecked", None, None),
 ("Anything that reads as machine-translated overall?", "", None, None)]
for i, r in enumerate(qs, 1):
    for j, v in enumerate(r, 1):
        c = q.cell(i, j, v); c.alignment = wrap
        c.font = hdr_font if i == 1 else Font(name=F, size=11)
        if i == 1: c.fill = hdr_fill
        elif j >= 3 and v is None: c.fill = inp
for j, w in enumerate([60, 22, 40, 40], 1): q.column_dimensions[chr(64 + j)].width = w

out = os.path.join(ROOT, "docs/i18n-review/translation-review.xlsx"); wb.save(out); print("wrote", out)
from collections import Counter
print(len(rows), Counter(r[0] for r in rows))
