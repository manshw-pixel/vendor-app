"""Apply a filled-in translation review back to web/src/i18n/{hi,mr}.json.

    python scripts/i18n/apply_review.py path/to/translation-review.xlsx [--dry-run]

Only rows with Verdict = "Change" and a non-empty "Your version" are applied. A row is
rejected (and reported) if its {{placeholders}} differ from the English, or if its key
no longer exists. Run the web tests afterwards: cd web && npx vitest run src/__tests__/i18n-*.
"""
import json, os, re, sys
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
I18N = os.path.join(ROOT, "web", "src", "i18n")
PH = re.compile(r"{{\s*\w+\s*}}")
ph = lambda s: sorted(p.replace(" ", "") for p in PH.findall(s))

def get(d, key):
    for part in key.split("."):
        if not isinstance(d, dict) or part not in d: return None
        d = d[part]
    return d

def put(raw, key, old, new):
    """Swap one value in the hand-formatted JSON text, leaving the layout untouched.
    Returns None when the `"leaf": "old"` pair is not unique -- apply that row by hand."""
    enc = lambda s: json.dumps(s, ensure_ascii=False)
    pair = f'{enc(key.split(".")[-1])}: {enc(old)}'
    if raw.count(pair) != 1: return None
    return raw.replace(pair, f'{enc(key.split(".")[-1])}: {enc(new)}')

def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) != 1: sys.exit(__doc__)
    dry = "--dry-run" in sys.argv
    wb = load_workbook(args[0], read_only=True)
    en = json.load(open(os.path.join(I18N, "en.json"), encoding="utf8"))
    rejected = 0
    for sheet, lang in (("Marathi", "mr"), ("Hindi", "hi")):
        path = os.path.join(I18N, f"{lang}.json")
        raw = open(path, encoding="utf8").read()
        data = json.loads(raw)
        applied = 0
        rows = wb[sheet].iter_rows(min_row=2, values_only=True)
        for _prio, _screen, key, _en, _cur, verdict, new, _comment in rows:
            if (verdict or "").strip().lower() != "change" or not (new or "").strip(): continue
            new, english = new.strip(), get(en, key)
            if not isinstance(english, str):
                print(f"REJECT {lang} {key}: key no longer exists"); rejected += 1; continue
            if ph(new) != ph(english):
                print(f"REJECT {lang} {key}: placeholders {ph(new)} != {ph(english)}"); rejected += 1; continue
            old = get(data, key)
            if old == new: continue
            swapped = put(raw, key, old, new)
            if swapped is None:
                print(f"REJECT {lang} {key}: value not unique in file, apply by hand"); rejected += 1; continue
            raw = swapped; applied += 1
        print(f"{lang}: {applied} change(s){' (dry run)' if dry else ''}")
        if applied and not dry:
            json.loads(raw)  # still valid JSON
            with open(path, "w", encoding="utf8", newline="") as f: f.write(raw)
    sys.exit(1 if rejected else 0)

if __name__ == "__main__": main()
